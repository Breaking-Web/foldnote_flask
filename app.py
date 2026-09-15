from flask import Flask, render_template, request, jsonify, send_from_directory
from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime
import json, uuid

app = Flask(__name__)
DATA = Path("data")
DATA.mkdir(exist_ok=True)
DB = DATA / "foldnote.xlsx"

def init_db():
    if DB.exists(): return
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws.append(["note_id","title","created_at","updated_at","tags"])
    seg = wb.create_sheet("Segments")
    seg.append(["segment_id","note_id","segment_order","original_text","created_at"])
    ver = wb.create_sheet("Versions")
    ver.append(["version_id","segment_id","version_order","text","label","created_at"])
    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "FoldNote – Översikt"
    dash["A3"] = "Öppna anteckningar och deras originalstycken samt alternativa versioner."
    wb.save(DB)

def rebuild_dashboard(book):
    if "Dashboard" in book.sheetnames:
        del book["Dashboard"]
    ws = book.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "FoldNote – Översikt"
    ws["A1"].font = __import__("openpyxl").styles.Font(size=22, bold=True, color="FFFFFF")
    ws["A1"].fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="111111")
    ws.merge_cells("A1:F1")
    ws["A2"] = "Automatiskt skapad översikt över alla anteckningar, stycken och alternativa versioner."
    ws.merge_cells("A2:F2")
    ws["A2"].font = __import__("openpyxl").styles.Font(italic=True, color="666666")
    widths = {"A":18,"B":28,"C":16,"D":55,"E":24,"F":65}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    row = 4
    notes = list(book["Notes"].iter_rows(min_row=2, values_only=True))
    segments = list(book["Segments"].iter_rows(min_row=2, values_only=True))
    versions = list(book["Versions"].iter_rows(min_row=2, values_only=True))
    seg_by_note = {}
    for s in segments: seg_by_note.setdefault(s[1], []).append(s)
    ver_by_seg = {}
    for v in versions: ver_by_seg.setdefault(v[1], []).append(v)
    for note in notes:
        ws.cell(row,1,"ANTECKNING")
        ws.cell(row,2,note[1])
        ws.cell(row,3,"Taggar")
        ws.cell(row,4,note[4] or "")
        ws.cell(row,5,"Uppdaterad")
        ws.cell(row,6,note[3] or "")
        for c in range(1,7):
            cell=ws.cell(row,c)
            cell.fill=__import__("openpyxl").styles.PatternFill("solid", fgColor="FF7A00")
            cell.font=__import__("openpyxl").styles.Font(bold=True, color="111111")
        row += 1
        ws.append(["Stycke #","Originaltext","Version #","Versionsnamn","Alternativ text",""])
        for c in range(1,6):
            ws.cell(row,c).font=__import__("openpyxl").styles.Font(bold=True,color="FFFFFF")
            ws.cell(row,c).fill=__import__("openpyxl").styles.PatternFill("solid",fgColor="222222")
        row += 1
        for idx,s in enumerate(sorted(seg_by_note.get(note[0],[]), key=lambda x:x[2]),1):
            vs = sorted(ver_by_seg.get(s[0],[]), key=lambda x:x[2])
            if not vs:
                ws.append([idx,s[3],"","",""])
                row += 1
            else:
                first=True
                for j,v in enumerate(vs,1):
                    ws.append([idx if first else "", s[3] if first else "", j, v[4], v[3]])
                    first=False; row += 1
        row += 2
    for r in ws.iter_rows():
        for cell in r:
            cell.alignment = __import__("openpyxl").styles.Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A4"

def wb():
    init_db()
    return load_workbook(DB)

def iso():
    return datetime.now().isoformat(timespec="seconds")

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/notes", methods=["GET"])
def notes():
    book=wb(); ws=book["Notes"]; out=[]
    for r in list(ws.iter_rows(min_row=2, values_only=True)):
        out.append({"id":r[0],"title":r[1],"created_at":r[2],"updated_at":r[3],"tags":r[4] or ""})
    return jsonify(out)

@app.route("/api/notes", methods=["POST"])
def create_note():
    d=request.json or {}; book=wb(); nid=str(uuid.uuid4()); now=iso()
    book["Notes"].append([nid,d.get("title","Ny anteckning"),now,now,d.get("tags","")])
    rebuild_dashboard(book)
    book.save(DB); return jsonify({"id":nid}),201

@app.route("/api/notes/<nid>", methods=["GET"])
def get_note(nid):
    book=wb(); nws=book["Notes"]; note=None
    for r in nws.iter_rows(min_row=2, values_only=True):
        if r[0]==nid: note={"id":r[0],"title":r[1],"created_at":r[2],"updated_at":r[3],"tags":r[4] or ""}; break
    if not note: return jsonify({"error":"not found"}),404
    segs=[]
    for r in book["Segments"].iter_rows(min_row=2, values_only=True):
        if r[1]==nid:
            versions=[]
            for v in book["Versions"].iter_rows(min_row=2, values_only=True):
                if v[1]==r[0]: versions.append({"id":v[0],"order":v[2],"text":v[3],"label":v[4]})
            versions.sort(key=lambda x:x["order"])
            segs.append({"id":r[0],"order":r[2],"original":r[3],"versions":versions})
    segs.sort(key=lambda x:x["order"]); note["segments"]=segs
    return jsonify(note)

@app.route("/api/notes/<nid>", methods=["PUT"])
def save_note(nid):
    d=request.json or {}; book=wb(); nws=book["Notes"]; found=False
    for row in range(2,nws.max_row+1):
        if nws.cell(row,1).value==nid:
            nws.cell(row,2).value=d.get("title","")
            nws.cell(row,4).value=iso()
            nws.cell(row,5).value=d.get("tags","")
            found=True; break
    if not found: return jsonify({"error":"not found"}),404
    # replace segments + versions for this note
    sws=book["Segments"]; vws=book["Versions"]
    old_ids={r[0] for r in sws.iter_rows(min_row=2, values_only=True) if r[1]==nid}
    for ws, colvals in [(vws, lambda r:r[1] in old_ids),(sws, lambda r:r[1]==nid)]:
        for row in range(ws.max_row,1,-1):
            vals=[ws.cell(row,c).value for c in range(1,ws.max_column+1)]
            if colvals(vals): ws.delete_rows(row)
    for i,s in enumerate(d.get("segments",[])):
        sid=s.get("id") or str(uuid.uuid4())
        sws.append([sid,nid,i,s.get("original",""),iso()])
        for j,v in enumerate(s.get("versions",[])):
            vws.append([v.get("id") or str(uuid.uuid4()),sid,j,v.get("text",""),v.get("label") or f"Version {j+1}",iso()])
    rebuild_dashboard(book)
    rebuild_dashboard(book)
    book.save(DB); return jsonify({"ok":True})

@app.route("/api/notes/<nid>", methods=["DELETE"])
def delete_note(nid):
    book=wb()
    sws=book["Segments"]; vws=book["Versions"]; nws=book["Notes"]
    sids={r[0] for r in sws.iter_rows(min_row=2, values_only=True) if r[1]==nid}
    for row in range(vws.max_row,1,-1):
        if vws.cell(row,2).value in sids: vws.delete_rows(row)
    for row in range(sws.max_row,1,-1):
        if sws.cell(row,2).value==nid: sws.delete_rows(row)
    for row in range(nws.max_row,1,-1):
        if nws.cell(row,1).value==nid: nws.delete_rows(row)
    rebuild_dashboard(book)
    rebuild_dashboard(book)
    book.save(DB); return jsonify({"ok":True})

@app.route("/api/export")
def export():
    init_db()
    book=wb(); rebuild_dashboard(book); book.save(DB)
    return send_from_directory(DATA, "foldnote.xlsx", as_attachment=True)

@app.route("/manifest.webmanifest")
def manifest():
    return send_from_directory("static","manifest.webmanifest", mimetype="application/manifest+json")

if __name__=="__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
